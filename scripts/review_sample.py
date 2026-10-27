# Copyright 2025 Julia Lenc
# Licensed under the Apache License, Version 2.0
# https://www.apache.org/licenses/LICENSE-2.0

"""
review_sample.py

Human review sample: measures how many LLM suggestions an Admin would
approve as they stand. The automatic check (evaluate_suggestions.py) can only
test what a rule can see; whether a suggestion is true, useful and worth
approving needs a person. Reviewing every row is the real workflow; this
script picks a smaller, stratified sample so the approval rate can be
measured in under an hour.

  make   Pick a sample from a review queue and write review_sample_<timestamp>.xlsx
         Tab A: 20 FLAGGED rows, spread evenly over R1-R5
         Tab B: 12 rows, spread over R6-R9 and the REVIEWED fields Prompt C rewrote
         Tab C: 8 REVIEWED fields Prompt C kept (do you agree they need no change?)
  score  Read a filled-in sample and print approval rates per tab and per rule
  rereview  Same rows as an earlier, filled-in sample, with the suggestions of a
         newer review queue. A row whose suggestion, Unverified and Missing are
         unchanged keeps its decision (Changed = no); the others are blank and
         show the previous suggestion and decision, so only they need reviewing

Decisions in the sample file:
  Tab A and Tab B   Approve (usable as written), Edit (right idea, needs changes),
                    Reject (wrong or not worth having)
  Tab C             Agree (keep is right), Disagree (it should have been rewritten)
  Note              optional; say why, e.g. "invented process", "too long"
  Unverified        claims the LLM said it could not confirm from the input
                    (check these first; blank for review queues made before Day 5)
  Missing           facts the LLM could not find in the input, as questions to the Admin

The sample file is named review_sample_*, so Script 2 never picks it up as a
review queue. The same seed always gives the same sample.

Usage:
  python scripts/review_sample.py make
      Samples the newest data/review_queue_*.xlsx.
  python scripts/review_sample.py make data/review_queue_20261025_135129.xlsx
  python scripts/review_sample.py score data/review_sample_20261025_190000.xlsx
  python scripts/review_sample.py rereview review_sample_20261025_202746_checked.xlsx
      Same rows, suggestions from the newest data/review_queue_*.xlsx
      (or pass the queue as a third argument).
"""

import json
import random
import sys
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT     = Path(__file__).resolve().parent.parent
DATA_DIR = ROOT / "data"
SEED     = 5

# Rows per tab. Tab B also covers the REVIEWED fields that Prompt C rewrote.
SIZES = {"A": 20, "B": 12, "C": 8}

COLUMNS   = ["Object", "Field_API_Name", "Field_Type", "Rule", "Original_Status",
             "Picklist_Values", "Current_Description", "LLM_Suggestion", "Unverified", "Missing", "LLM_Reasoning",
             "Decision", "Note"]
DECISIONS = {"A": ["Approve", "Edit", "Reject"],
             "B": ["Approve", "Edit", "Reject"],
             "C": ["Agree", "Disagree"]}
WIDTHS    = {"Object": 12, "Field_API_Name": 30, "Field_Type": 14, "Rule": 8,
             "Original_Status": 13, "Picklist_Values": 25, "Current_Description": 50, "LLM_Suggestion": 50,
             "Unverified": 35, "Missing": 35, "LLM_Reasoning": 40, "Decision": 12, "Note": 30}
SHEETS    = {"A": "Tab A — FLAGGED", "B": "Tab B — UNCERTAIN", "C": "Tab C — Kept"}

# Extra columns of a re-review sample
REREVIEW_COLUMNS = ["Changed", "Previous_Suggestion", "Previous_Decision", "Previous_Note"]
WIDTHS.update({"Changed": 10, "Previous_Suggestion": 50, "Previous_Decision": 12, "Previous_Note": 30})


def read_tab(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h or "") for h in rows[0]]
    return [dict(zip(headers, r)) for r in rows[1:] if any(v not in (None, "") for v in r)]


def find_sheet(wb, prefix: str):
    for name in wb.sheetnames:
        if name.startswith(prefix):
            return wb[name]
    sys.exit(f"Sheet starting with '{prefix}' not found. Sheets: {wb.sheetnames}")


def stratum(row: dict) -> str:
    """Group a row by its rule; Prompt C rewrites have no rule and form their own group."""
    if row.get("Original_Status") == "REVIEWED":
        return "C-rewrite"
    return row.get("Rule") or "-"


def spread(rows: list[dict], n: int, rng: random.Random) -> list[dict]:
    """Take n rows, round-robin over the strata so every rule is represented."""
    groups = defaultdict(list)
    for r in rows:
        groups[stratum(r)].append(r)
    for g in groups.values():
        rng.shuffle(g)
    picked, keys = [], sorted(groups)
    while len(picked) < n and any(groups.values()):
        for k in keys:
            if groups[k] and len(picked) < n:
                picked.append(groups[k].pop())
    return sorted(picked, key=lambda r: (stratum(r), r.get("Object") or "", r.get("Field_API_Name") or ""))


def read_queue(source: Path) -> dict[str, list[dict]]:
    wb_in = openpyxl.load_workbook(source, read_only=True)
    # Tab C of the review queue has no suggestion column: a kept field's suggestion is its original
    return {"A": read_tab(find_sheet(wb_in, "Tab A")),
            "B": read_tab(find_sheet(wb_in, "Tab B")),
            "C": [dict(r, Rule="", LLM_Suggestion=r.get("Current_Description"))
                  for r in read_tab(find_sheet(wb_in, "Tab C")) if r.get("Original_Status") == "REVIEWED"]}


def add_picklists(sample: dict[str, list[dict]]) -> None:
    """The review queue has no picklist values; the reviewer needs them to check value names."""
    picklists = {}
    classified = DATA_DIR / "sf_classified.json"
    if classified.exists():
        for f in json.loads(classified.read_text(encoding="utf-8")).get("fields", []):
            picklists[(f.get("object"), f.get("field_api_name"))] = ", ".join(f.get("picklist_values") or [])
    for rows in sample.values():
        for r in rows:
            r["Picklist_Values"] = picklists.get((r.get("Object"), r.get("Field_API_Name")), "")


def write_sample(sample: dict[str, list[dict]], columns: list[str], keep_decisions: bool = False) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for tab, rows in sample.items():
        ws = wb.create_sheet(SHEETS[tab])
        ws.append(columns)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
        for r in rows:
            ws.append([r.get(c, "") if keep_decisions or c not in ("Decision", "Note") else "" for c in columns])
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        for ci, h in enumerate(columns, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(ci)].width = WIDTHS[h]
        ws.freeze_panes = "A2"
        dv = DataValidation(type="list", formula1='"' + ",".join(DECISIONS[tab]) + '"', allow_blank=True)
        col = openpyxl.utils.get_column_letter(columns.index("Decision") + 1)
        dv.sqref = f"{col}2:{col}{len(rows) + 1}"
        ws.add_data_validation(dv)

    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    out = DATA_DIR / f"review_sample_{timestamp}.xlsx"
    wb.save(out)
    return out


def make(source: Path) -> Path:
    queue  = read_queue(source)
    rng    = random.Random(SEED)
    sample = {tab: spread(queue[tab], SIZES[tab], rng) for tab in ("A", "B", "C")}
    add_picklists(sample)
    out = write_sample(sample, COLUMNS)
    print(f"Sample from {source.name}: Tab A {len(sample['A'])}, Tab B {len(sample['B'])}, "
          f"Tab C {len(sample['C'])} rows -> {out.name}")
    return out


def _norm(value) -> str:
    return " ".join(str(value or "").split())


def rereview(previous: Path, source: Path) -> Path:
    """
    Same rows as a filled-in sample, suggestions from a newer queue. A field
    stays on the tab it was sampled on, so approval rates stay comparable.
    """
    queue = read_queue(source)
    by_key = {}
    for tab in ("C", "B", "A"):          # A and B win over C if a field is on both
        for r in queue[tab]:
            by_key[(r.get("Object"), r.get("Field_API_Name"))] = r

    wb_prev = openpyxl.load_workbook(previous, read_only=True)
    sample, changed, missing_rows = {}, 0, []
    for tab in ("A", "B", "C"):
        rows = []
        for old in read_tab(find_sheet(wb_prev, f"Tab {tab}")):
            key = (old.get("Object"), old.get("Field_API_Name"))
            new = by_key.get(key)
            if new is None:
                missing_rows.append(f"{key[0]}.{key[1]}")
                continue
            row = {c: new.get(c, old.get(c, "")) for c in COLUMNS}
            same = all(_norm(new.get(c)) == _norm(old.get(c)) for c in ("LLM_Suggestion", "Unverified", "Missing"))
            row.update(Changed="no" if same else "yes",
                       Previous_Suggestion=old.get("LLM_Suggestion"),
                       Previous_Decision=old.get("Decision"), Previous_Note=old.get("Note"),
                       Decision=old.get("Decision") if same else "", Note=old.get("Note") if same else "")
            changed += not same
            rows.append(row)
        sample[tab] = rows
    add_picklists(sample)
    out = write_sample(sample, COLUMNS + REREVIEW_COLUMNS, keep_decisions=True)
    total = sum(len(r) for r in sample.values())
    print(f"Re-review of {previous.name} with {source.name}: {changed} of {total} rows changed "
          f"and need a decision; {total - changed} kept their decision -> {out.name}")
    if missing_rows:
        print(f"Not in the new queue (left out): {', '.join(missing_rows)}")
    return out


def pct(n: int, d: int) -> str:
    return f"{n}/{d} ({100 * n / d:.1f}%)" if d else "0/0"


def score(path: Path) -> int:
    wb = openpyxl.load_workbook(path, read_only=True)
    blank = 0
    for tab in ("A", "B", "C"):
        rows    = read_tab(find_sheet(wb, f"Tab {tab}"))
        first   = DECISIONS[tab][0]
        by_rule = defaultdict(lambda: defaultdict(int))
        for r in rows:
            d = (r.get("Decision") or "").strip().capitalize()
            if d not in DECISIONS[tab]:
                blank += 1
                d = "(blank)"
            by_rule[stratum(r) if tab != "C" else "kept"][d] += 1
            by_rule["all"][d] += 1
        print(f"\n{SHEETS[tab]}  ({len(rows)} rows)")
        for rule in sorted(k for k in by_rule if k != "all") + ["all"]:
            counts = by_rule[rule]
            total  = sum(counts.values())
            detail = ", ".join(f"{d} {counts[d]}" for d in DECISIONS[tab] + ["(blank)"] if counts[d])
            print(f"  {rule:<10} {first}: {pct(counts[first], total):<14}  {detail}")
        notes = [f"{r.get('Field_API_Name')}: {r.get('Note')}" for r in rows if r.get("Note")]
        for n in notes:
            print(f"    note  {n}")

    a_rows = read_tab(find_sheet(wb, "Tab A"))
    a_ok   = sum((r.get("Decision") or "").strip().capitalize() == "Approve" for r in a_rows)
    print(f"\nDefinition of Done (wiki 03): Tab A approvable without editing >= 80%: "
          f"{pct(a_ok, len(a_rows))} -> {'met' if a_rows and a_ok / len(a_rows) >= 0.8 else 'not met'}")
    if blank:
        print(f"{blank} row(s) have no valid decision")
    return 0


def newest_queue() -> Path:
    queues = sorted(DATA_DIR.glob("review_queue_*.xlsx"))
    if not queues:
        sys.exit("No data/review_queue_*.xlsx found. Run Script 1 first.")
    return queues[-1]


def main() -> int:
    if len(sys.argv) < 2 or sys.argv[1] not in ("make", "score", "rereview"):
        print(__doc__)
        return 1
    if sys.argv[1] == "make":
        make(Path(sys.argv[2]) if len(sys.argv) > 2 else newest_queue())
        return 0
    if sys.argv[1] == "rereview":
        if len(sys.argv) < 3:
            sys.exit("rereview needs the filled-in sample file")
        rereview(Path(sys.argv[2]), Path(sys.argv[3]) if len(sys.argv) > 3 else newest_queue())
        return 0
    if len(sys.argv) < 3:
        sys.exit("score needs the filled-in sample file")
    return score(Path(sys.argv[2]))


if __name__ == "__main__":
    sys.exit(main())
